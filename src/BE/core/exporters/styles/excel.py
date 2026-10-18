from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Palette & Colors
COLOR_TITLE = "111827"
COLOR_SUBTITLE = "6B7280"
COLOR_CARD_BORDER = "D1D5DB"
COLOR_TABLE_BORDER = "D1D5DB"
COLOR_INNER_GRID = "E5E7EB"

COLOR_CARD_FOLLOWING = "2563EB"
COLOR_CARD_FOLLOWERS = "16A34A"
COLOR_CARD_NON_FOLLOWERS = "DC2626"
COLOR_CARD_NEW_UNFOLLOWERS = "D97706"

COLOR_HEADER_BG = "F9FAFB"
COLOR_HEADER_TEXT = "374151"
COLOR_ROW_BG_WHITE = "FFFFFF"
COLOR_ROW_BG_LIGHT = "F9FAFB"
COLOR_TEXT_PRIMARY = "374151"
COLOR_TEXT_SECONDARY = "4B5563"
COLOR_LINK = "2563EB"

# Layout & Dimensions
ROW_INDEX_TITLE = 1
ROW_INDEX_SUBTITLE = 2
ROW_INDEX_SPACER_1 = 3
ROW_INDEX_CARD_VAL = 4
ROW_INDEX_CARD_LBL = 5
ROW_INDEX_SPACER_2 = 6
ROW_INDEX_HEADER = 7
ROW_INDEX_DATA_START = 8

ROW_HEIGHT_TITLE = 28
ROW_HEIGHT_SUBTITLE = 18
ROW_HEIGHT_SPACER_1 = 10
ROW_HEIGHT_CARD_VAL = 28
ROW_HEIGHT_CARD_LBL = 18
ROW_HEIGHT_SPACER_2 = 14
ROW_HEIGHT_HEADER = 24
ROW_HEIGHT_DATA = 20

COL_WIDTH_INDEX = 12
COL_WIDTH_USERNAME = 26
COL_WIDTH_URL = 46
COL_WIDTH_DATE = 26

COL_WIDTHS = {
                    1: COL_WIDTH_INDEX,
                    2: COL_WIDTH_USERNAME,
                    3: COL_WIDTH_URL,
                    4: COL_WIDTH_DATE,
                }

# Content & Labels
REPORT_TITLE = "Who is the Mole? The Full Non-Followers Report"
DEFAULT_SHEET_TITLE = "Non-Followers"
DEFAULT_REPORT_FILENAME = "non_followers_report"
TABLE_HEADERS = ["#", "Username", "Profile URL", "Unfollow Date"]

# Typography & Fonts
FONT_FAMILY = "Calibri"

FONT_TITLE = Font(name = FONT_FAMILY, size = 16, bold = True, color = COLOR_TITLE)
FONT_SUBTITLE = Font(name = FONT_FAMILY, size = 9, color = COLOR_SUBTITLE)
FONT_CARD_LABEL = Font(name = FONT_FAMILY, size = 9, bold = True, color = COLOR_SUBTITLE)
FONT_HEADER = Font(name = FONT_FAMILY, size = 10, bold = True, color = COLOR_HEADER_TEXT)
FONT_INDEX = Font(name = FONT_FAMILY, size = 10, color = COLOR_TEXT_SECONDARY)
FONT_USERNAME = Font(name = FONT_FAMILY, size = 10, color = COLOR_TEXT_PRIMARY)
FONT_URL = Font(name = FONT_FAMILY, size = 10, color = COLOR_LINK, underline = "single")
FONT_URL_DEFAULT = Font(name = FONT_FAMILY, size = 10, color = COLOR_TEXT_PRIMARY)
FONT_DATE = Font(name = FONT_FAMILY, size = 10, color = COLOR_TEXT_SECONDARY)

# Borders
SIDE_CARD = Side(style = 'thin', color = COLOR_CARD_BORDER)
SIDE_TABLE = Side(style = 'thin', color = COLOR_TABLE_BORDER)
SIDE_INNER = Side(style = 'thin', color = COLOR_INNER_GRID)

BORDER_CARD_TOP = Border(top = SIDE_CARD, left = SIDE_CARD, right = SIDE_CARD)
BORDER_CARD_BOTTOM = Border(bottom = SIDE_CARD, left = SIDE_CARD, right = SIDE_CARD)
BORDER_HEADER = Border(top = SIDE_TABLE, bottom = SIDE_TABLE, left = SIDE_TABLE, right = SIDE_TABLE)
BORDER_DATA_CELL = Border(top = SIDE_INNER, bottom = SIDE_INNER, left = SIDE_INNER, right = SIDE_INNER)

# Fills
FILL_WHITE = PatternFill(start_color = COLOR_ROW_BG_WHITE, end_color = COLOR_ROW_BG_WHITE, fill_type = "solid")
FILL_LIGHT = PatternFill(start_color = COLOR_ROW_BG_LIGHT, end_color = COLOR_ROW_BG_LIGHT, fill_type = "solid")
FILL_HEADER = PatternFill(start_color = COLOR_HEADER_BG, end_color = COLOR_HEADER_BG, fill_type = "solid")

# Alignments
ALIGN_BANNER = Alignment(horizontal = "left", vertical = "center", wrap_text = False)
ALIGN_CENTER = Alignment(horizontal = "center", vertical = "center", wrap_text = True)
ALIGN_LEFT = Alignment(horizontal = "left", vertical = "center", wrap_text = True)