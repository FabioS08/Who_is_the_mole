from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl import Workbook

from ..utils import ensure_file_extension

from rich.console import Console

from typing import List, Dict


console = Console()


# Palette
COLOR_SLATE_DARK = "1E293B"
COLOR_SLATE_LIGHT = "F8FAFC"
COLOR_WHITE = "FFFFFF"
COLOR_BORDER = "CBD5E1"
COLOR_TEXT_PRIMARY = "334155"
COLOR_TEXT_SECONDARY = "64748B"
COLOR_LINK = "2563EB"
COLOR_ALERT = "DC2626"

# Layout & Dimensions
ROW_INDEX_TITLE = 1
ROW_INDEX_COUNT = 2
ROW_INDEX_SPACER = 3
ROW_INDEX_HEADER = 4

ROW_HEIGHT_TITLE = 28
ROW_HEIGHT_COUNT = 20
ROW_HEIGHT_SPACER = 12
ROW_HEIGHT_HEADER = 26
ROW_HEIGHT_DATA = 20

COL_WIDTH_USERNAME = 25
COL_WIDTH_URL = 45
COL_WIDTH_DATE = 22

COL_WIDTHS = {
                1: COL_WIDTH_USERNAME,
                2: COL_WIDTH_URL,
                3: COL_WIDTH_DATE
            }

# Content & Labels
REPORT_TITLE = "Non-Followers Report"
DEFAULT_SHEET_TITLE = "Non-Followers"
DEFAULT_REPORT_FILENAME = "non_followers_report"
TABLE_HEADERS = ["Username", "Profile URL", "Unfollow Date"]

# Typography & Fonts
FONT_FAMILY = "Calibri"

FONT_TITLE = Font(name = FONT_FAMILY, size = 16, bold = True, color = COLOR_SLATE_DARK)
FONT_COUNT = Font(name = FONT_FAMILY, size = 11, bold = True, color = COLOR_ALERT)
HEADER_FONT = Font(name = FONT_FAMILY, size = 11, bold = True, color = COLOR_WHITE)
FONT_USERNAME = Font(name = FONT_FAMILY, size = 10, bold = True, color = COLOR_TEXT_PRIMARY)
FONT_URL = Font(name = FONT_FAMILY, size = 10, color = COLOR_LINK, underline = "single")
FONT_URL_DEFAULT = Font(name = FONT_FAMILY, size = 10, color = COLOR_TEXT_PRIMARY)
FONT_DATE = Font(name = FONT_FAMILY, size = 10, color = COLOR_TEXT_SECONDARY)

# Styles
HEADER_FILL = PatternFill(start_color = COLOR_SLATE_DARK, end_color = COLOR_SLATE_DARK, fill_type = "solid")
ROW_BG_WHITE = PatternFill(start_color = COLOR_WHITE, end_color = COLOR_WHITE, fill_type = "solid")
ROW_BG_LIGHT = PatternFill(start_color = COLOR_SLATE_LIGHT, end_color = COLOR_SLATE_LIGHT, fill_type = "solid")

BORDER_THIN = Border(
                        left = Side(style = 'thin', color = COLOR_BORDER),
                        right = Side(style = 'thin', color = COLOR_BORDER),
                        top = Side(style = 'thin', color = COLOR_BORDER),
                        bottom = Side(style = 'thin', color = COLOR_BORDER)
                    )

ALIGN_CENTER = Alignment(horizontal = "center", vertical = "center", wrap_text = True)
ALIGN_LEFT = Alignment(horizontal = "left", vertical = "center", wrap_text = True)
ALIGN_BANNER = Alignment(horizontal = "left", vertical = "center", wrap_text = False)


def _build_header(ws: Worksheet, total_count: int) -> int:

    '''
    It sets up the title banner and total count summary rows in the worksheet.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     total_count:int
      The total number of detected non-followers

    Outputs
    -------
     next_row:int
      The starting row index for the table headers
    '''

    ws.row_dimensions[ROW_INDEX_TITLE].height = ROW_HEIGHT_TITLE
    cell_title = ws.cell(row = ROW_INDEX_TITLE, column = 1, value = REPORT_TITLE)
    cell_title.font = FONT_TITLE
    cell_title.alignment = ALIGN_BANNER

    ws.row_dimensions[ROW_INDEX_COUNT].height = ROW_HEIGHT_COUNT
    cell_count = ws.cell(row = ROW_INDEX_COUNT, column = 1, value = f"Total Non-Followers: {total_count}")
    cell_count.font = FONT_COUNT
    cell_count.alignment = ALIGN_BANNER

    ws.row_dimensions[ROW_INDEX_SPACER].height = ROW_HEIGHT_SPACER

    return ROW_INDEX_HEADER


def _build_table_headers(ws: Worksheet, row: int) -> None:

    '''
    It creates and styles the column headers row for the non-followers table.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     row:int
      The row number where headers should be placed

    Outputs
    -------
     None
    '''

    ws.row_dimensions[row].height = ROW_HEIGHT_HEADER

    for col_idx, header_text in enumerate(TABLE_HEADERS, start = 1):

        cell = ws.cell(row = row, column = col_idx, value = header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN


def _render_username_cell(ws: Worksheet, row: int, username: str, fill: PatternFill) -> None:

    '''
    It creates and formats the username cell with bold slate styling and thin border.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     row:int
      The row index where the cell should be placed

     username:str
      The target account username

     fill:PatternFill
      The background fill pattern for alternating rows

    Outputs
    -------
     None
    '''

    cell = ws.cell(row = row, column = 1, value = username)
    cell.font = FONT_USERNAME
    cell.alignment = ALIGN_LEFT
    cell.fill = fill
    cell.border = BORDER_THIN


def _render_url_cell(ws: Worksheet, row: int, url: str, fill: PatternFill) -> None:

    '''
    It creates and formats the profile URL cell, configuring hyperlinks for web URLs.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     row:int
      The row index where the cell should be placed

     url:str
      The profile URL address

     fill:PatternFill
      The background fill pattern for alternating rows

    Outputs
    -------
     None
    '''

    cell = ws.cell(row = row, column = 2, value = url)
    cell.alignment = ALIGN_LEFT
    cell.fill = fill
    cell.border = BORDER_THIN

    if url and (str(url).startswith("http://") or str(url).startswith("https://")):
        cell.hyperlink = url
        cell.font = FONT_URL

    else:
        cell.font = FONT_URL_DEFAULT


def _render_date_cell(ws: Worksheet, row: int, date_str: str, fill: PatternFill) -> None:

    '''
    It creates and formats the unfollow date cell with centered alignment and subtle text color.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     row:int
      The row index where the cell should be placed

     date_str:str
      The recorded unfollow date or history status string

     fill:PatternFill
      The background fill pattern for alternating rows

    Outputs
    -------
     None
    '''

    cell = ws.cell(row = row, column = 3, value = date_str)
    cell.font = FONT_DATE
    cell.alignment = ALIGN_CENTER
    cell.fill = fill
    cell.border = BORDER_THIN


def _populate_row(ws: Worksheet, row: int, item: Dict, is_even: bool) -> None:

    '''
    It formats and populates a single data row with username, profile URL and unfollow date.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     row:int
      The row number to populate

     item:Dict
      The dictionary containing non-follower user details

     is_even:bool
      Whether the current row index is even, for alternating background selection

    Outputs
    -------
     None
    '''

    ws.row_dimensions[row].height = ROW_HEIGHT_DATA
    row_fill = ROW_BG_LIGHT if is_even else ROW_BG_WHITE

    _render_username_cell(ws, row, item.get('username', 'N/A'), row_fill)
    _render_url_cell(ws, row, item.get('url', 'N/A'), row_fill)
    _render_date_cell(ws, row, item.get('unfollow_date', 'N/A'), row_fill)


def _populate_data_rows(ws: Worksheet, data: List[Dict], start_row: int) -> None:

    '''
    It populates the worksheet with formatted user data rows, alternating backgrounds and hyperlinks.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

     data:List[Dict]
      The list of non-followers data dictionaries

     start_row:int
      The starting row number for data rows

    Outputs
    -------
     None
    '''

    for idx, item in enumerate(data):
        
        current_row = start_row + idx
        _populate_row(ws, current_row, item, is_even = (idx % 2 == 1))


def _apply_column_widths(ws: Worksheet) -> None:

    '''
    It sets predefined column widths for the table columns.

    Parameters
    ----------
     ws:Worksheet
      The active openpyxl worksheet

    Outputs
    -------
     None
    '''

    for col_idx, width in COL_WIDTHS.items():

        col_letter = get_column_letter(col_idx)
        col_dim = ws.column_dimensions[col_letter]

        if col_dim is not None:
            col_dim.width = float(width)


def _setup_workbook(title: str = DEFAULT_SHEET_TITLE) -> tuple[Workbook, Worksheet]:

    '''
    It initializes the workbook and configures the active worksheet with title and view settings.

    Parameters
    ----------
     title:str
      The title to assign to the main worksheet (default: 'Non-Followers')

    Outputs
    -------
     tuple[Workbook, Worksheet]
      The initialized Workbook instance and its configured active Worksheet
    '''

    wb = Workbook()
    ws = wb.active

    if ws is None:
        ws = wb.create_sheet(title = title)
    else:
        ws.title = title

    # Ensure grid lines are not visible
    ws.views.sheetView[0].showGridLines = False

    return wb, ws


def export_to_excel(data: List[Dict], filename: str = DEFAULT_REPORT_FILENAME) -> None:

    '''
    It generates an Excel spreadsheet report containing the non-followers list in a styled table.

    Parameters
    ----------
     data:List[Dict]
      The list of dictionaries containing user details and unfollow dates

     filename:str
      The output Excel file path or name (default: 'non_followers_report')

    Outputs
    -------
     None
    '''

    filename = ensure_file_extension(filename, ".xlsx")

    wb, ws = _setup_workbook(DEFAULT_SHEET_TITLE)

    header_row = _build_header(ws, len(data))
    _build_table_headers(ws, header_row)
    _populate_data_rows(ws, data, header_row + 1)
    _apply_column_widths(ws)

    # Freeze panes below table header
    ws.freeze_panes = f"A{header_row + 1}"

    wb.save(filename)
    console.print(f"[bold green]✓[/] Report successfully saved to [cyan]{filename}[/cyan]")